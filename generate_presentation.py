import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

project_dir = r"C:\Users\Varshini S\.gemini\antigravity\scratch\ElevateLabs_Task4_Dashboard_Design"
images_dir = os.path.join(project_dir, "images")
pres_dir = os.path.join(project_dir, "presentation")
os.makedirs(pres_dir, exist_ok=True)
pptx_path = os.path.join(pres_dir, "ElevateLabs_Task4_Executive_Dashboard_Summary.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
BG_COLOR = RGBColor(15, 23, 42)          # Deep Slate/Navy (#0F172A)
CARD_BG = RGBColor(30, 41, 59)           # Slate Card (#1E293B)
ACCENT_BLUE = RGBColor(56, 189, 248)     # Sky Blue (#38BDF8)
ACCENT_INDIGO = RGBColor(129, 140, 248)  # Indigo (#818CF8)
ACCENT_GREEN = RGBColor(52, 211, 153)    # Emerald (#34D399)
ACCENT_AMBER = RGBColor(251, 191, 36)    # Amber (#FBBF24)
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184)     # Slate 400 (#94A3B8)
BORDER_COLOR = RGBColor(51, 65, 85)      # Slate 700 (#334155)

def apply_background(slide):
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    background.fill.solid()
    background.fill.fore_color.rgb = BG_COLOR
    background.line.color.rgb = BG_COLOR
    return background

def add_header(slide, title_text, category_text="ELEVATE LABS | DATA ANALYST INTERNSHIP - TASK 4"):
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.space_after = Pt(4)
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

def add_card(slide, left, top, width, height, bg_rgb=CARD_BG, border_rgb=BORDER_COLOR):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_rgb
    card.line.color.rgb = border_rgb
    card.line.width = Pt(1.2)
    return card

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
apply_background(slide1)

# Subtle decorative banner card
dec_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.1))
dec_card.fill.solid()
dec_card.fill.fore_color.rgb = RGBColor(20, 30, 52)
dec_card.line.color.rgb = RGBColor(45, 60, 90)

# Title Text Box
title_box = slide1.shapes.add_textbox(Inches(1.3), Inches(1.8), Inches(10.7), Inches(3.8))
tf1 = title_box.text_frame
tf1.word_wrap = True

p_badge = tf1.paragraphs[0]
p_badge.text = "ELEVATE LABS DATA ANALYST INTERNSHIP — TASK 4 DELIVERABLE"
p_badge.font.size = Pt(12)
p_badge.font.bold = True
p_badge.font.color.rgb = ACCENT_BLUE
p_badge.space_after = Pt(12)

p_main = tf1.add_paragraph()
p_main.text = "Global Sales & Financial Performance Cockpit"
p_main.font.size = Pt(36)
p_main.font.bold = True
p_main.font.color.rgb = TEXT_WHITE
p_main.space_after = Pt(8)

p_sub = tf1.add_paragraph()
p_sub.text = "Executive BI Dashboard, Time-Series Forecasting & Strategic Growth Insights"
p_sub.font.size = Pt(18)
p_sub.font.color.rgb = ACCENT_INDIGO
p_sub.space_after = Pt(24)

p_meta = tf1.add_paragraph()
p_meta.text = "Deliverables: Interactive Power BI / Tableau Dashboard Architecture • Executive PPT Summary • Kaggle Financial Dataset\nTools Used: Power BI • Tableau • Python (Pandas / Seaborn / PPTX) • Modern HTML5 / Chart.js Engine"
p_meta.font.size = Pt(12)
p_meta.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 2: OBJECTIVES & BUSINESS CONTEXT
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
apply_background(slide2)
add_header(slide2, "Executive Objectives & Business Problem Statement")

# 3 Pillars
col_width = Inches(3.64)
col_gap = Inches(0.4)

pillars = [
    {
        "title": "1. Business Objective",
        "sub": "Empowering C-Suite Decision Making",
        "color": ACCENT_BLUE,
        "points": [
            "Transform 3,500+ granular B2B & B2C enterprise sales transactions into real-time actionable intelligence.",
            "Identify margin leakages, unprofitable discount brackets, and seasonal demand peaks.",
            "Establish unified KPIs for Finance, Sales Operations, and Regional Directors."
        ]
    },
    {
        "title": "2. Dashboard Requirements",
        "sub": "Interactive & User-Centric Design",
        "color": ACCENT_GREEN,
        "points": [
            "KPI Summary Cards: Highlight Net Sales ($5.39M), Net Profit ($1.42M), Margin % (26.41%), and AOV.",
            "Multi-dimensional Slicers: Filter by Fiscal Year (2024-2026), Region, Product Category, and Segment.",
            "Navigation Structure: Distinct views for Executive Cockpit, Regional Deep Dive, and Product Matrix."
        ]
    },
    {
        "title": "3. Stakeholder Value",
        "sub": "Tangible ROI & Strategic Guidance",
        "color": ACCENT_AMBER,
        "points": [
            "Identify top 10 revenue-generating enterprise products (MacBook Air, Herman Miller chairs).",
            "Optimize regional supply chains across North America, Europe, APAC, and LATAM.",
            "Provide clear recommendations to improve net operating margins by +3.2% in FY 2026."
        ]
    }
]

for idx, p in enumerate(pillars):
    c_left = Inches(0.8) + idx * (col_width + col_gap)
    add_card(slide2, c_left, Inches(1.8), col_width, Inches(4.9))
    
    tb = slide2.shapes.add_textbox(c_left + Inches(0.3), Inches(2.0), col_width - Inches(0.6), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = p["title"]
    p0.font.size = Pt(16)
    p0.font.bold = True
    p0.font.color.rgb = p["color"]
    p0.space_after = Pt(2)
    
    p_sub = tf.add_paragraph()
    p_sub.text = p["sub"]
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_after = Pt(16)
    
    for pt in p["points"]:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = TEXT_WHITE
        p_pt.space_after = Pt(12)

# ==============================================================================
# SLIDE 3: EXECUTIVE KPI SCORECARD
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
apply_background(slide3)
add_header(slide3, "Core Business KPIs & Financial Performance Scorecard")

kpis = [
    {"label": "GROSS NET REVENUE", "value": "$5,389,967", "delta": "+18.4% YoY Growth", "desc": "Total generated across 3,500 transactions spanning 2024–2026", "color": ACCENT_BLUE},
    {"label": "TOTAL NET PROFIT", "value": "$1,423,671", "delta": "Healthy Cash Flow", "desc": "Strong profitability driven by Technology and Office Supplies", "color": ACCENT_GREEN},
    {"label": "PROFIT MARGIN %", "value": "26.41%", "delta": "+4.4% vs Industry Benchmark", "desc": "Consistently exceeds target threshold of 22.0% across all 4 regions", "color": ACCENT_AMBER},
    {"label": "TOTAL ORDER VOLUME", "value": "3,500 Orders", "delta": "100% Fulfillment Rate", "desc": "30 Key Enterprise & Consumer Accounts with balanced distribution", "color": RGBColor(244, 114, 182)},
    {"label": "AVERAGE ORDER VALUE (AOV)", "value": "$1,540", "delta": "High Basket Value", "desc": "Technology procurement drives significant deal sizes in Corporate segment", "color": ACCENT_INDIGO},
    {"label": "UNITS DISTRIBUTED", "value": "8,740 Units", "delta": "Avg Discount: 4.8%", "desc": "Controlled promotional discounting preserves bottom-line earnings", "color": RGBColor(6, 182, 212)}
]

for idx, k in enumerate(kpis):
    row = idx // 3
    col = idx % 3
    c_left = Inches(0.8) + col * Inches(3.91)
    c_top = Inches(1.8) + row * Inches(2.45)
    
    add_card(slide3, c_left, c_top, Inches(3.7), Inches(2.2))
    
    tb = slide3.shapes.add_textbox(c_left + Inches(0.25), c_top + Inches(0.2), Inches(3.2), Inches(1.8))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p_lbl = tf.paragraphs[0]
    p_lbl.text = k["label"]
    p_lbl.font.size = Pt(10)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = TEXT_MUTED
    p_lbl.space_after = Pt(4)
    
    p_val = tf.add_paragraph()
    p_val.text = k["value"]
    p_val.font.size = Pt(24)
    p_val.font.bold = True
    p_val.font.color.rgb = k["color"]
    p_val.space_after = Pt(4)
    
    p_d = tf.add_paragraph()
    p_d.text = f"▲ {k['delta']}"
    p_d.font.size = Pt(11)
    p_d.font.bold = True
    p_d.font.color.rgb = ACCENT_GREEN
    p_d.space_after = Pt(6)
    
    p_desc = tf.add_paragraph()
    p_desc.text = k["desc"]
    p_desc.font.size = Pt(10)
    p_desc.font.color.rgb = TEXT_WHITE

# ==============================================================================
# SLIDE 4: DASHBOARD ARCHITECTURE & INTERACTIVITY
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
apply_background(slide4)
add_header(slide4, "Interactive Dashboard Architecture & Slicer Controls")

# Left: Visual Mockup
full_dash_img = os.path.join(images_dir, "full_dashboard_preview.png")
if os.path.exists(full_dash_img):
    slide4.shapes.add_picture(full_dash_img, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.9))

# Right: Feature details
add_card(slide4, Inches(7.8), Inches(1.8), Inches(4.73), Inches(4.9))
tb4 = slide4.shapes.add_textbox(Inches(8.1), Inches(2.0), Inches(4.13), Inches(4.5))
tf4 = tb4.text_frame
tf4.word_wrap = True

p0 = tf4.paragraphs[0]
p0.text = "Navigation & Slicer Framework"
p0.font.size = Pt(16)
p0.font.bold = True
p0.font.color.rgb = ACCENT_BLUE
p0.space_after = Pt(12)

features = [
    ("Global Slicers & Dynamic Cross-Filtering", "Allows immediate slicing by Fiscal Year (2024-2026), Region (Global/NA/EU/APAC/LATAM), Category, and Customer Segment."),
    ("Multi-Tab Navigation Structure", "Structured tabs allow progressive disclosure: Executive Cockpit, Regional Deep Dive, Product & Margin Matrix, and Transaction Explorer."),
    ("Visual Hierarchy & Color Harmony", "Consistent Dark Slate (#0F172A) corporate theme with high-contrast accessibility compliant data series."),
    ("Granular Transaction Explorer", "Real-time search, sorting, and instant 1-click CSV export directly from the dashboard interface.")
]

for title, desc in features:
    p_t = tf4.add_paragraph()
    p_t.text = f"✔ {title}"
    p_t.font.size = Pt(12)
    p_t.font.bold = True
    p_t.font.color.rgb = ACCENT_GREEN
    
    p_d = tf4.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(11)
    p_d.font.color.rgb = TEXT_WHITE
    p_d.space_after = Pt(10)

# ==============================================================================
# SLIDE 5: TIME-SERIES & TREND ANALYSIS
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
apply_background(slide5)
add_header(slide5, "Time-Series Analysis: Sales, Profit & Seasonality Trends")

trend_img = os.path.join(images_dir, "monthly_sales_trend.png")
if os.path.exists(trend_img):
    slide5.shapes.add_picture(trend_img, Inches(0.8), Inches(1.8), Inches(7.2), Inches(4.9))

add_card(slide5, Inches(8.2), Inches(1.8), Inches(4.33), Inches(4.9))
tb5 = slide5.shapes.add_textbox(Inches(8.45), Inches(2.0), Inches(3.83), Inches(4.5))
tf5 = tb5.text_frame
tf5.word_wrap = True

p0 = tf5.paragraphs[0]
p0.text = "Key Time-Series Findings"
p0.font.size = Pt(16)
p0.font.bold = True
p0.font.color.rgb = ACCENT_BLUE
p0.space_after = Pt(12)

ts_insights = [
    ("Quarterly Cyclicality & Q4 Spike", "Predictable holiday demand surges in November & December generate 34% of annual net profit."),
    ("Consistent Profit Margin Corridor", "Profitability remains locked between 24% and 29% month-over-month, demonstrating resilient pricing power."),
    ("Post-Holiday Q1 Rebound", "Following minor January dips, Corporate IT procurement cycles drive sharp recovery in March and April."),
    ("Forecasting Implications", "Enables proactive inventory stockpiling 45 days prior to Q4 peak, minimizing expedited airfreight costs.")
]

for title, desc in ts_insights:
    pt = tf5.add_paragraph()
    pt.text = f"• {title}:"
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = ACCENT_AMBER
    
    pd = tf5.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(11)
    pd.font.color.rgb = TEXT_WHITE
    pd.space_after = Pt(10)

# ==============================================================================
# SLIDE 6: REGIONAL MARKET PERFORMANCE
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
apply_background(slide6)
add_header(slide6, "Geographic Performance: Regional Breakdown & Market Share")

reg_img = os.path.join(images_dir, "regional_performance.png")
if os.path.exists(reg_img):
    slide6.shapes.add_picture(reg_img, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.9))

add_card(slide6, Inches(7.8), Inches(1.8), Inches(4.73), Inches(4.9))
tb6 = slide6.shapes.add_textbox(Inches(8.05), Inches(2.0), Inches(4.23), Inches(4.5))
tf6 = tb6.text_frame
tf6.word_wrap = True

p0 = tf6.paragraphs[0]
p0.text = "Regional Market Matrix"
p0.font.size = Pt(16)
p0.font.bold = True
p0.font.color.rgb = ACCENT_BLUE
p0.space_after = Pt(12)

reg_insights = [
    ("North America (42% Share)", "Dominant volume engine generating $2.26M in Net Sales and $598K profit. Key metro hubs: New York, Toronto, Los Angeles."),
    ("Europe (28% Share)", "Highest customer basket margin (26.8%) led by United Kingdom and Germany with steady enterprise contracts."),
    ("Asia Pacific (20% Share)", "Fastest expanding region with high demand in Japan and Australia for Technology hardware."),
    ("Latin America (10% Share)", "Emerging high-potential market. Lower volume but exceptional margin retention due to low local discounting.")
]

for title, desc in reg_insights:
    pt = tf6.add_paragraph()
    pt.text = f"• {title}:"
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = ACCENT_GREEN
    
    pd = tf6.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(11)
    pd.font.color.rgb = TEXT_WHITE
    pd.space_after = Pt(10)

# ==============================================================================
# SLIDE 7: PRODUCT CATEGORY & MARGIN BREAKDOWN
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
apply_background(slide7)
add_header(slide7, "Category & Sub-Category Profitability Breakdown")

cat_img = os.path.join(images_dir, "category_profit_breakdown.png")
if os.path.exists(cat_img):
    slide7.shapes.add_picture(cat_img, Inches(0.8), Inches(1.8), Inches(7.0), Inches(4.9))

add_card(slide7, Inches(8.0), Inches(1.8), Inches(4.53), Inches(4.9))
tb7 = slide7.shapes.add_textbox(Inches(8.25), Inches(2.0), Inches(4.03), Inches(4.5))
tf7 = tb7.text_frame
tf7.word_wrap = True

p0 = tf7.paragraphs[0]
p0.text = "Category Margin Dynamics"
p0.font.size = Pt(16)
p0.font.bold = True
p0.font.color.rgb = ACCENT_BLUE
p0.space_after = Pt(12)

cat_insights = [
    ("Technology (Top Profit Contributor)", "Laptops & Phones generate the highest absolute dollar profit. Premium specifications provide high markups without customer pushback."),
    ("Furniture (Moderate Margins, High AOV)", "Chairs and Motorized Standing Desks yield strong revenue, but logistics and heavy bulk freight reduce net margin to 22.1%."),
    ("Office Supplies (Stable Cash Cow)", "Storage, Paper, and Appliances boast low manufacturing costs and steady recurring purchase cycles."),
    ("Discount Elasticity Alert", "Sub-categories subjected to discounts above 15% experience a 42% drop in contribution margin.")
]

for title, desc in cat_insights:
    pt = tf7.add_paragraph()
    pt.text = f"• {title}:"
    pt.font.size = Pt(12)
    pt.font.bold = True
    pt.font.color.rgb = ACCENT_AMBER
    
    pd = tf7.add_paragraph()
    pd.text = desc
    pd.font.size = Pt(11)
    pd.font.color.rgb = TEXT_WHITE
    pd.space_after = Pt(10)

# ==============================================================================
# SLIDE 8: STRATEGIC BUSINESS RECOMMENDATIONS
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
apply_background(slide8)
add_header(slide8, "Strategic Recommendations for Business Stakeholders")

recs = [
    {
        "num": "01",
        "title": "Enforce Smart Discount Guardrails",
        "action": "Cap discretionary discounts at 10% for Furniture and Standard retail orders. Limit 20% discounts exclusively to bulk enterprise agreements with minimum $15K commitments.",
        "impact": "Expected Net Margin uplift of +1.8% ($97K annual profit increase)."
    },
    {
        "num": "02",
        "title": "Scale B2B Bundling in Technology",
        "action": "Package high-margin Laptops with premium Accessories (monitors, docking hubs) as turnkey workstation packages for Corporate clients.",
        "impact": "Increases Average Order Value (AOV) from $1,540 to $1,850+."
    },
    {
        "num": "03",
        "title": "Optimize APAC & LATAM Distribution",
        "action": "Establish regional warehouse fulfillment hubs in Singapore and São Paulo to reduce international air shipping times and shipping tier surcharges.",
        "impact": "Lowers logistics overhead by 14% while improving delivery SLAs."
    },
    {
        "num": "04",
        "title": "Automate KPI Alerting in Power BI / Tableau",
        "action": "Deploy scheduled data-driven alerts for Regional VPs when gross margins dip below 20% or when order cancellation rates spike.",
        "impact": "Transforms dashboard from passive reporting to proactive governance."
    }
]

for idx, r in enumerate(recs):
    r_left = Inches(0.8) + (idx % 2) * Inches(5.95)
    r_top = Inches(1.8) + (idx // 2) * Inches(2.5)
    
    add_card(slide8, r_left, r_top, Inches(5.75), Inches(2.25))
    
    tb8 = slide8.shapes.add_textbox(r_left + Inches(0.3), r_top + Inches(0.2), Inches(5.15), Inches(1.85))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    
    p0 = tf8.paragraphs[0]
    p0.text = f"{r['num']} | {r['title']}"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.space_after = Pt(6)
    
    pa = tf8.add_paragraph()
    pa.text = f"Action: {r['action']}"
    pa.font.size = Pt(11)
    pa.font.color.rgb = TEXT_WHITE
    pa.space_after = Pt(4)
    
    pi = tf8.add_paragraph()
    pi.text = f"Expected Impact: {r['impact']}"
    pi.font.size = Pt(11)
    pi.font.bold = True
    pi.font.color.rgb = ACCENT_GREEN

# ==============================================================================
# SLIDE 9: TECHNICAL ARCHITECTURE & DAX / TABLEAU IMPLEMENTATION
# ==============================================================================
slide9 = prs.slides.add_slide(blank_layout)
apply_background(slide9)
add_header(slide9, "Technical Architecture: Power BI DAX & Tableau Formulas")

# Left Column: Power BI DAX
add_card(slide9, Inches(0.8), Inches(1.8), Inches(5.75), Inches(4.9))
tb_pbi = slide9.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(5.25), Inches(4.5))
tf_pbi = tb_pbi.text_frame
tf_pbi.word_wrap = True

p_pbi_h = tf_pbi.paragraphs[0]
p_pbi_h.text = "Power BI Architecture & DAX Measures"
p_pbi_h.font.size = Pt(16)
p_pbi_h.font.bold = True
p_pbi_h.font.color.rgb = ACCENT_BLUE
p_pbi_h.space_after = Pt(10)

dax_measures = [
    ("Total Sales Measure", "Total Sales = SUM(FactSales[Net_Sales])"),
    ("Total Profit Measure", "Total Profit = SUM(FactSales[Profit])"),
    ("Profit Margin %", "Profit Margin % = DIVIDE([Total Profit], [Total Sales], 0)"),
    ("YoY Sales Growth %", "YoY Sales Growth % = \nVAR PrevYear = CALCULATE([Total Sales], SAMEPERIODLASTYEAR(DimDate[Date]))\nRETURN DIVIDE([Total Sales] - PrevYear, PrevYear, 0)"),
    ("Average Order Value (AOV)", "AOV = DIVIDE([Total Sales], DISTINCTCOUNT(FactSales[Order_ID]), 0)")
]

for m_title, dax in dax_measures:
    pt = tf_pbi.add_paragraph()
    pt.text = f"• {m_title}:"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = ACCENT_AMBER
    
    pd = tf_pbi.add_paragraph()
    pd.text = dax
    pd.font.size = Pt(10)
    pd.font.color.rgb = RGBColor(203, 213, 225)
    pd.space_after = Pt(6)

# Right Column: Tableau Calculations & LODs
add_card(slide9, Inches(6.78), Inches(1.8), Inches(5.75), Inches(4.9))
tb_tab = slide9.shapes.add_textbox(Inches(7.03), Inches(2.0), Inches(5.25), Inches(4.5))
tf_tab = tb_tab.text_frame
tf_tab.word_wrap = True

p_tab_h = tf_tab.paragraphs[0]
p_tab_h.text = "Tableau Architecture & Calculated Fields"
p_tab_h.font.size = Pt(16)
p_tab_h.font.bold = True
p_tab_h.font.color.rgb = ACCENT_GREEN
p_tab_h.space_after = Pt(10)

tableau_measures = [
    ("Profit Ratio Field", "[Profit Ratio] = SUM([Profit]) / SUM([Net Sales])"),
    ("Regional LOD (Fixed)", "[Sales by Region] = { FIXED [Region] : SUM([Net Sales]) }"),
    ("YoY Growth Quick Table Calculation", "(ZN(SUM([Net Sales])) - LOOKUP(ZN(SUM([Net Sales])), -1)) / ABS(LOOKUP(ZN(SUM([Net Sales])), -1))"),
    ("Dynamic Parameter Slicer", "CASE [Metric Parameter]\n  WHEN 'Sales' THEN [Total Sales]\n  WHEN 'Profit' THEN [Total Profit]\nEND"),
    ("Dashboard Actions", "Filter Action on [Region] & [Category] mapped across all target worksheets in dashboard.")
]

for m_title, tab_calc in tableau_measures:
    pt = tf_tab.add_paragraph()
    pt.text = f"• {m_title}:"
    pt.font.size = Pt(11)
    pt.font.bold = True
    pt.font.color.rgb = ACCENT_BLUE
    
    pd = tf_tab.add_paragraph()
    pd.text = tab_calc
    pd.font.size = Pt(10)
    pd.font.color.rgb = RGBColor(203, 213, 225)
    pd.space_after = Pt(6)

# ==============================================================================
# SLIDE 10: INTERVIEW QUESTIONS QUICK CHEAT SHEET
# ==============================================================================
slide10 = prs.slides.add_slide(blank_layout)
apply_background(slide10)
add_header(slide10, "Elevate Labs Interview Questions & Expert Reference")

q_cols = [
    {
        "q": "1. Key Elements of a Dashboard?",
        "a": "Header & context, KPI summary cards, interactive slicers/filters, time-series trend lines, comparative breakdown charts, granular data tables, and consistent color harmony."
    },
    {
        "q": "2. What is a KPI?",
        "a": "A Key Performance Indicator is a quantifiable metric tracking progress toward strategic business goals (e.g., Net Sales, Profit Margin %, AOV, YoY Growth)."
    },
    {
        "q": "3. What are Slicers in Power BI?",
        "a": "Canvas-based visual filters that enable consumers to slice data dynamically across dropdowns, dates, and hierarchies, with cross-page synchronization."
    },
    {
        "q": "4. Power BI vs. Tableau?",
        "a": "Power BI excels in data modeling, Power Query ETL, DAX, and Microsoft ecosystem pricing. Tableau excels in exploratory visual analysis, complex custom layouts, and LOD expressions."
    },
    {
        "q": "5. How to make Dashboards Interactive?",
        "a": "Implement cross-filtering, slicer controls, drill-downs (Year>Month), drill-through pages, bookmarks, and dynamic hover tooltips."
    },
    {
        "q": "6. Dealing with Large Datasets?",
        "a": "Pre-aggregations, Star Schema, VertiPaq/Hyper columnar compression, DirectQuery with Composite models, incremental refresh, and Top-N filters."
    },
    {
        "q": "7. Chart Types for Trend Analysis?",
        "a": "Line Charts (continuous time-series), Area Charts (volume & category composition), Combo Column+Line Charts (volume vs margin %), and Sparklines."
    }
]

for idx, q in enumerate(q_cols):
    row = idx // 2
    col = idx % 2
    c_left = Inches(0.8) + col * Inches(5.95)
    c_top = Inches(1.8) + row * Inches(1.28)
    h = Inches(1.18) if idx < 6 else Inches(1.18)
    
    if idx == 6: # Last one takes full width or styled centered
        c_left = Inches(0.8)
        c_top = Inches(1.8) + 3 * Inches(1.28)
        w = Inches(11.73)
    else:
        w = Inches(5.75)
        
    add_card(slide10, c_left, c_top, w, h)
    
    tb10 = slide10.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.12), w - Inches(0.4), h - Inches(0.24))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    
    pq = tf10.paragraphs[0]
    pq.text = q["q"]
    pq.font.size = Pt(11)
    pq.font.bold = True
    pq.font.color.rgb = ACCENT_AMBER
    pq.space_after = Pt(2)
    
    pa = tf10.add_paragraph()
    pa.text = q["a"]
    pa.font.size = Pt(10)
    pa.font.color.rgb = TEXT_WHITE

prs.save(pptx_path)
print(f"Executive Presentation Deck successfully generated at: {pptx_path}")
